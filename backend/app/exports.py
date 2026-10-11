"""Export mapped reports with exact decimal text and escaped account labels."""
import io
from xml.sax.saxutils import escape

def workbook(entity,statement,trial):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book=Workbook();book.remove(book.active)
    sections=[('Trial Balance',['Code','Account','Category','Debit','Credit','Balance'],[[r['code'],r['name'],r['category'],r['debit'],r['credit'],r['balance']] for r in trial]),('Profit and Loss',['Account','Category','Amount'],[[r['name'],r['category'],r['amount']] for r in statement['profit_and_loss']]),('Balance Sheet',['Account','Category','Amount'],[[r['name'],r['category'],r['amount']] for r in statement['balance_sheet']])]
    for name,headers,rows in sections:
        sheet=book.create_sheet(name);sheet.append([entity.name,entity.currency,statement['period']]);sheet.append(['Amounts are exact decimal text; verify source completeness before relying on statements.']);sheet.append(headers)
        for row in rows:
            sheet.append(row)
            # Explicit text cells prevent formula execution and large-number rounding.
            for cell in sheet[sheet.max_row]: cell.data_type='s'
        for cell in sheet[3]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='142C4B')
        sheet.freeze_panes='A4';sheet.auto_filter.ref=f'A3:{sheet.cell(3,len(headers)).column_letter}{sheet.max_row}'
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(50,max(18,max(len(str(c.value or '')) for c in col)+2))
    summary=book.create_sheet('Source and Controls');summary.append(['Basis',statement['basis']]);summary.append(['Balance check',statement['balance_check']]);summary.append(['Unclosed earnings',statement['unclosed_earnings']])
    summary.append(['Source journal IDs']);
    for identifier in statement['journal_ids']: summary.append([identifier])
    for row in summary:
        for cell in row:
            if cell.value is not None: cell.data_type='s'
    summary.column_dimensions['A'].width=45;summary.column_dimensions['B'].width=90
    for row in summary:
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
    buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()

def pdf_report(entity,statement):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    buffer=io.BytesIO();styles=getSampleStyleSheet();styles['BodyText'].fontSize=9
    doc=SimpleDocTemplate(buffer,pagesize=A4,rightMargin=40,leftMargin=40,topMargin=40,bottomMargin=40)
    story=[Paragraph('AI Finance OS — Financial reporting snapshot',styles['Title']),Paragraph(escape(entity.name)+' · '+escape(entity.currency)+' · '+statement['period'],styles['Heading2']),Paragraph(escape(statement['basis']),styles['BodyText']),Spacer(1,18)]
    for title,rows in [('Profit and loss',statement['profit_and_loss']),('Balance sheet',statement['balance_sheet'])]:
        story.append(Paragraph(title,styles['Heading2']))
        table_data=[['Account','Classification','Amount']]+[[Paragraph(escape(r['name']),styles['BodyText']),r['category'],r['amount']] for r in rows]
        table=Table(table_data,colWidths=[220,100,115],repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#142c4b')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),8),('VALIGN',(0,0),(-1,-1),'TOP'),('RIGHTPADDING',(0,0),(-1,-1),8),('LEFTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#dce3ed')),('ALIGN',(2,1),(2,-1),'RIGHT')]))
        story.extend([table,Spacer(1,16)])
    for key in ('revenue','expenses','net_income','assets','liabilities','equity','unclosed_earnings','balance_check'):
        story.append(Paragraph(key.replace('_',' ').title()+': '+statement[key],styles['BodyText']))
    story.append(Spacer(1,16));story.append(Paragraph(f"Source journals: {len(statement['journal_ids'])}. This snapshot is based on imported records and explicit mappings; it is not an assurance opinion.",styles['BodyText']))
    def footer(canvas,document):
        canvas.saveState();canvas.setFont('Helvetica',8);canvas.drawString(40,25,'AI Finance OS · Review source completeness and opening balances');canvas.drawRightString(555,25,f'Page {document.page}');canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer);return buffer.getvalue()
