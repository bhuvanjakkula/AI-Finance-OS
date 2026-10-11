import io, uuid
from decimal import Decimal
from openpyxl import load_workbook
from test_platform import workspace, journal, ingest, member

def test_versioned_budget_actual_comparison(workspace):
    c,h,e,_,_=workspace
    ingest(c,h,e,[journal(value='120')])
    body={'entity_id':e,'period':'2026-10','account_code':'4000','amount':'100','version':1}
    assert c.post('/budgets',headers=h,json=body).status_code==201
    first=c.get('/reports/budget',headers=h,params={'entity_id':e,'period':'2026-10'}).json()['rows'][0]
    assert first['actual']=='120.00' and first['variance']=='20.00'
    body['amount']='110';body['version']=2
    assert c.post('/budgets',headers=h,json=body).status_code==201
    latest=c.get('/reports/budget',headers=h,params={'entity_id':e,'period':'2026-10'}).json()['rows'][0]
    assert latest['budget_version']==2 and latest['variance']=='10.00'
    assert c.post('/budgets',headers=h,json=body).status_code==422

def test_intercompany_independent_elimination_and_stale_snapshot(workspace):
    c,h,e,_,slug=workspace
    other=c.post('/entities',headers=h,json={'code':'EU','name':'Europe','currency':'USD'}).json()['id']
    for entity,code,category in [(e,'1100','asset'),(other,'2100','liability'),(other,'5000','expense')]:
        assert c.post(f'/entities/{entity}/accounts',headers=h,json={'code':code,'name':code,'category':category,'group_code':code}).status_code==201
    sale={'external_id':'IC-SALE','posting_date':'2026-10-01','currency':'USD','lines':[{'account_code':'1100','debit':'100','credit':'0','reference':'IC:EU:GROUP-1'},{'account_code':'4000','debit':'0','credit':'100'}]}
    bill={'external_id':'IC-BILL','posting_date':'2026-10-01','currency':'USD','lines':[{'account_code':'5000','debit':'100','credit':'0'},{'account_code':'2100','debit':'0','credit':'100','reference':'IC:US:GROUP-1'}]}
    assert ingest(c,h,e,[sale]).status_code==200 and ingest(c,h,other,[bill]).status_code==200
    body={'entity_ids':[e,other],'period':'2026-10','group_currency':'USD','rates':{}}
    match=c.post('/intercompany/runs',headers=h,json=body);assert match.status_code==200,match.text
    assert len(match.json()['matches'])==1
    proposal=c.post('/consolidation/elimination-proposals',headers=h,json={'source_run_id':match.json()['id'],'reason':'Eliminate reciprocal balances'}).json()['id']
    url='/consolidation/elimination-proposals/'+proposal+'/decision';decision={'decision':'approved','note':'Verified source documents'}
    assert c.post(url,headers=h,json=decision).status_code==403
    reviewer,_=member(c,h,slug)
    assert c.post(url,headers=reviewer,json=decision).status_code==200
    body['approved_elimination_ids']=[proposal]
    consolidated=c.post('/consolidation/runs',headers=h,json=body);assert consolidated.status_code==200,consolidated.text
    rows={r['group_code']:r['balance'] for r in consolidated.json()['rows']}
    assert rows['1100']=='0.00' and rows['2100']=='0.00'
    assert consolidated.json()['applied_elimination_ids']==[proposal]
    assert ingest(c,h,e,[journal('NEW')]).status_code==200
    assert c.post('/consolidation/runs',headers=h,json=body).status_code==409

def test_excel_pdf_exports_and_source_snapshot(workspace):
    c,h,e,_,_=workspace
    assert ingest(c,h,e,[journal(value='123.45')]).status_code==200
    query={'entity_id':e,'period':'2026-10','format':'xlsx'}
    excel=c.get('/reports/export',headers=h,params=query);assert excel.status_code==200,excel.text[:200]
    book=load_workbook(io.BytesIO(excel.content),data_only=False)
    assert set(book.sheetnames)=={'Trial Balance','Profit and Loss','Balance Sheet','Source and Controls'}
    assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
    query['format']='pdf';pdf=c.get('/reports/export',headers=h,params=query)
    assert pdf.status_code==200 and pdf.content.startswith(b'%PDF-')
    snapshots=[r for r in c.get('/runs',headers=h).json() if r['kind']=='report_snapshot']
    assert len(snapshots)==2 and snapshots[0]['output']['statement']['journal_ids']

def test_transaction_date_is_separate_from_posting_period(workspace):
    c,h,e,_,_=workspace
    entry=journal();entry['transaction_date']='2026-09-28'
    assert ingest(c,h,e,[entry]).status_code==200
    row=c.get('/journals',headers=h,params={'entity_id':e}).json()[0]
    assert row['transaction_date']=='2026-09-28' and row['posting_date']=='2026-10-01' and row['period']=='2026-10'
